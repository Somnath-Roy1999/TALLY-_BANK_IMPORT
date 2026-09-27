import tkinter as tk
from tkinter import ttk,filedialog,messagebox
from pathlib import Path
from datetime import datetime
import xml.etree.ElementTree as ET
import traceback

try:
    import openpyxl
except ImportError:
    root=tk.Tk(); root.withdraw()
    messagebox.showerror("ACFC Tally Voucher Generator", "Python package 'openpyxl' is missing.\\n\\nOpen Command Prompt and run:\\n\\npip install openpyxl\\n\\nThen start the program again.")
    root.destroy()
    raise

APP="ACFC Tally Voucher Generator"; COMPANY="ACFC E SERVICES INDIA PRIVATE LIMITED"
BASE_REQ=["Date","Narration","ledger type"]

def s(x): return str(x or "").strip()
def k(x): return s(x).casefold()
def amt(x):
    if x in (None,""): return 0.0
    return float(str(x).replace(",","").strip())
def date(x):
    if isinstance(x,datetime): return x.strftime("%Y%m%d")
    for f in ("%d/%m/%y","%d/%m/%Y","%Y-%m-%d","%d-%m-%Y","%m/%d/%Y"):
        try:return datetime.strptime(str(x).strip(),f).strftime("%Y%m%d")
        except:pass
    if isinstance(x,(int,float)):
        from openpyxl.utils.datetime import from_excel
        return from_excel(x).strftime("%Y%m%d")
    raise ValueError("Invalid date: "+repr(x))
def add(p,t,v=None,**a):
    e=ET.SubElement(p,t,a)
    if v is not None:e.text=str(v)
    return e

def master(path):
    w=openpyxl.load_workbook(path,data_only=True); ws=w["MASTER"] if "MASTER" in w.sheetnames else w.active
    h={k(ws.cell(1,c).value):c for c in range(1,ws.max_column+1)}
    if "ledger name" not in h or "under" not in h: raise ValueError("Master must contain Ledger Name and Under columns.")
    m={}
    for r in range(2,ws.max_row+1):
        n=s(ws.cell(r,h["ledger name"]).value); u=s(ws.cell(r,h["under"]).value)
        if n:m[k(n)]=(n,u or "Sundry Creditors")
    return m
def rows(path):
    w=openpyxl.load_workbook(path,data_only=True); ws=w.active
    h={k(ws.cell(1,c).value):c for c in range(1,ws.max_column+1)}

    # Only columns required by the voucher types actually present are checked.
    def has(name): return k(name) in h
    def value(row,name):
        col=h.get(k(name))
        return ws.cell(row,col).value if col else None

    missing=[x for x in BASE_REQ if not has(x)]
    if missing:
        raise ValueError("Bank Excel missing: "+", ".join(missing))

    types=set()
    for r in range(2,ws.max_row+1):
        raw=value(r,"ledger type")
        if raw not in (None,""):
            types.add(k(raw))

    type_requirements={
        "journal":["Journal Ledger","Journal Payment Bank"],
        "payment":["Payment bank"],
        "contra":["Contra Credit bank","Contra Debit Bank"],
        "receipt":["Receipt Bank"],
        "salary":["Payment bank"],
    }
    for typ, reqs in type_requirements.items():
        if typ in types:
            missing.extend(x for x in reqs if not has(x))

    # Amount can be in Withdrawal Amt., Deposit Amt., or either one depending
    # on the source statement. At least one amount column must exist.
    if not has("Withdrawal Amt.") and not has("Deposit Amt."):
        missing.append("Withdrawal Amt. or Deposit Amt.")

    if missing:
        raise ValueError("Bank Excel missing: "+", ".join(dict.fromkeys(missing)))

    out=[]
    for r in range(2,ws.max_row+1):
        z={
            "Date":value(r,"Date"),
            "Narration":value(r,"Narration"),
            "Withdrawal Amt.":value(r,"Withdrawal Amt."),
            "Deposit Amt.":value(r,"Deposit Amt."),
            "ledger type":value(r,"ledger type"),
            "Ledger Name":value(r,"Ledger Name"),
            "Journal Ledger":value(r,"Journal Ledger"),
            "CONTRA BANK":value(r,"CONTRA BANK"),
            "Payment bank":value(r,"Payment bank"),
            "Contra Credit bank":value(r,"Contra Credit bank"),
            "Contra Debit Bank":value(r,"Contra Debit Bank"),
            "Journal Payment Bank":value(r,"Journal Payment Bank"),
            "Receipt Bank":value(r,"Receipt Bank"),
        }
        if any(v not in (None,"") for v in z.values()):
            out.append(z)
    return out

def build(rs,m,auto=True):
    vs=[]; missing=[]; cnt={"Journal":0,"Payment":0,"Contra":0,"Receipt":0}; total=0
    q={x:0 for x in cnt}

    def ensure(n,u):
        n=s(n)
        if n and k(n) not in m:
            missing.append((n,u))
            if auto:m[k(n)]=(n,u)

    def amount(r):
        # Prefer withdrawal, then deposit. This matches bank-statement layouts
        # while also supporting files where only one amount column exists.
        return amt(r.get("Withdrawal Amt.")) or amt(r.get("Deposit Amt."))

    for i,r in enumerate(rs,2):
        t=k(r.get("ledger type"))
        d=date(r.get("Date")); nar=s(r.get("Narration"))
        led=s(r.get("Ledger Name"))

        if t=="journal":
            a=amount(r); jl=s(r.get("Journal Ledger")); bank=s(r.get("Journal Payment Bank"))
            if not jl or not bank:
                raise ValueError(f"Row {i}: Journal needs Journal Ledger and Journal Payment Bank.")
            if not led:
                raise ValueError(f"Row {i}: Journal needs Ledger Name.")
            ensure(jl,"Indirect Expenses"); ensure(led,"Sundry Creditors"); ensure(bank,"Bank Accounts")
            if a<=0: raise ValueError(f"Row {i}: Journal amount invalid.")
            q["Journal"]+=1; q["Payment"]+=1; cnt["Journal"]+=1; cnt["Payment"]+=1; total+=a
            vs += [
                (d,"Journal",f"J-{q['Journal']:04d}",nar,[(jl,-a,1),(led,a,0)]),
                (d,"Payment",f"P-{q['Payment']:04d}",nar,[(led,-a,1),(bank,a,0)])
            ]

        elif t=="payment":
            a=amount(r); bank=s(r.get("Payment bank"))
            if not led:
                raise ValueError(f"Row {i}: Payment needs Ledger Name.")
            if not bank:
                raise ValueError(f"Row {i}: Payment needs Payment bank.")
            ensure(led,"Sundry Creditors"); ensure(bank,"Bank Accounts")
            if a<=0: raise ValueError(f"Row {i}: Payment amount invalid.")
            q["Payment"]+=1; cnt["Payment"]+=1; total+=a
            vs.append((d,"Payment",f"P-{q['Payment']:04d}",nar,[(led,-a,1),(bank,a,0)]))

        elif t=="contra":
            a=amount(r)
            credit=s(r.get("Contra Credit bank"))
            debit=s(r.get("Contra Debit Bank"))
            if not credit or not debit:
                raise ValueError(f"Row {i}: Contra needs Contra Credit bank and Contra Debit Bank.")
            ensure(debit,"Bank Accounts"); ensure(credit,"Bank Accounts")
            if a<=0: raise ValueError(f"Row {i}: Contra amount invalid.")
            q["Contra"]+=1; cnt["Contra"]+=1; total+=a
            # Contra Debit Bank = DR, Contra Credit bank = CR.
            vs.append((d,"Contra",f"C-{q['Contra']:04d}",nar,[(debit,-a,1),(credit,a,0)]))

        elif t in ("salary","salary payment","salarypayment"):
            # Salary entries are handled as Payment vouchers.
            a=amount(r); bank=s(r.get("Payment bank"))
            if not led:
                raise ValueError(f"Row {i}: Salary needs Ledger Name.")
            if not bank:
                raise ValueError(f"Row {i}: Salary needs Payment bank.")
            ensure(led,"Sundry Creditors"); ensure(bank,"Bank Accounts")
            if a<=0: raise ValueError(f"Row {i}: Salary amount invalid.")
            q["Payment"]+=1; cnt["Payment"]+=1; total+=a
            vs.append((d,"Payment",f"P-{q['Payment']:04d}",nar,[(led,-a,1),(bank,a,0)]))

        elif t=="receipt":
            a=amount(r); bank=s(r.get("Receipt Bank"))
            if not led:
                raise ValueError(f"Row {i}: Receipt needs Ledger Name.")
            if not bank:
                raise ValueError(f"Row {i}: Receipt needs Receipt Bank.")
            ensure(bank,"Bank Accounts"); ensure(led,"Sundry Creditors")
            if a<=0: raise ValueError(f"Row {i}: Receipt amount invalid.")
            q["Receipt"]+=1; cnt["Receipt"]+=1; total+=a
            vs.append((d,"Receipt",f"R-{q['Receipt']:04d}",nar,[(led,a,0),(bank,-a,1)]))

        elif t=="":
            raise ValueError(f"Row {i}: Ledger type blank.")
        else:
            raise ValueError(f"Row {i}: Unsupported ledger type {r.get('ledger type')!r}.")

    # Remove duplicate missing-ledger notices while preserving order.
    missing=list(dict.fromkeys(missing))
    return vs,cnt,total,missing

def voucher_xml(vs,path):
    e=ET.Element("ENVELOPE");h=ET.SubElement(e,"HEADER")
    for t,v in [("VERSION","1"),("TALLYREQUEST","Import"),("TYPE","Data"),("ID","Vouchers")]:add(h,t,v)
    b=ET.SubElement(e,"BODY");d=ET.SubElement(b,"DESC");sv=ET.SubElement(d,"STATICVARIABLES");add(sv,"SVCURRENTCOMPANY",COMPANY);add(sv,"SVEXPORTFORMAT","$$SysName:XML");data=ET.SubElement(b,"DATA")
    for dt,typ,no,nar,ls in vs:
        v=ET.SubElement(ET.SubElement(data,"TALLYMESSAGE"),"VOUCHER",{"VCHTYPE":typ,"ACTION":"Create","OBJVIEW":"Accounting Voucher View"})
        for t,x in [("DATE",dt),("EFFECTIVEDATE",dt),("VOUCHERTYPENAME",typ),("VOUCHERNUMBER",no),("REFERENCE",no),("VCHENTRYMODE","Accounting Voucher"),("ISINVOICE","No"),("PERSISTEDVIEW","Accounting Voucher View"),("NARRATION",nar)]:add(v,t,x)
        for j,(ln,a,dr) in enumerate(ls):
            z=ET.SubElement(v,"ALLLEDGERENTRIES.LIST");add(z,"LEDGERNAME",ln);add(z,"ISDEEMEDPOSITIVE","Yes" if dr else "No");add(z,"ISLASTDEEMEDPOSITIVE","Yes" if j==len(ls)-1 else "No");add(z,"AMOUNT",f"{a:.2f}")
    ET.ElementTree(e).write(path,encoding="utf-8",xml_declaration=True)
def master_xml(m,path):
    e=ET.Element("ENVELOPE");h=ET.SubElement(e,"HEADER")
    for t,v in [("VERSION","1"),("TALLYREQUEST","Import"),("TYPE","Data"),("ID","All Masters")]:add(h,t,v)
    b=ET.SubElement(e,"BODY");d=ET.SubElement(b,"DESC");sv=ET.SubElement(d,"STATICVARIABLES");add(sv,"SVCURRENTCOMPANY",COMPANY);add(sv,"SVEXPORTFORMAT","$$SysName:XML");data=ET.SubElement(b,"DATA")
    for n,u in sorted(m.values(),key=lambda x:k(x[0])):
        l=ET.SubElement(ET.SubElement(data,"TALLYMESSAGE"),"LEDGER",{"NAME":n,"ACTION":"Create"});add(l,"NAME",n);add(l,"PARENT",u);add(l,"ISBILLWISEON","Yes" if k(u)=="sundry creditors" else "No");add(l,"AFFECTSSTOCK","No")
        if k(u)=="bank accounts":add(l,"ISBANKINGENABLED","Yes")
    ET.ElementTree(e).write(path,encoding="utf-8",xml_declaration=True)
def review(vs,path):
    w=openpyxl.Workbook();ws=w.active;ws.title="Voucher Review";ws.append(["Date","Voucher Type","Voucher No","Narration","Ledger Name","Under","DR","CR"])
    for dt,t,no,nar,ls in vs:
        for ln,a,dr in ls:ws.append([datetime.strptime(dt,"%Y%m%d").date(),t,no,nar,ln,"",abs(a) if dr else "",abs(a) if not dr else ""])
    ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:c.font=openpyxl.styles.Font(bold=True)
    for c,wid in zip("ABCDEFGH",[13,14,14,60,32,24,16,16]):ws.column_dimensions[c].width=wid
    w.save(path)
BUILTIN_MASTER = [("TIS EXPENSES","Indirect Expenses"),("BISWA BARAN DUBE","Sundry Creditors"),("MD PERSONAL A/C","Bank Accounts"),("HDFC BANK 3231","Bank Accounts"),("ACFC E SERVICES INDIA PRIVATE LIMITED","Sundry Creditors")]
def builtin_master(): return {k(n):(n,u) for n,u in BUILTIN_MASTER}
def generate(bank,out,auto=True):
    if not s(bank):
        raise ValueError("Please select a Bank Statement Excel file.")
    bank_path=Path(bank).expanduser().resolve()
    if not bank_path.is_file():
        raise ValueError(f"Bank Statement file not found: {bank_path}")
    # If Output Folder is left blank, use the bank statement's folder.
    # This avoids accidentally writing to an unrelated current working directory.
    out_path=Path(out).expanduser() if s(out) else bank_path.parent
    out_path=out_path.resolve()
    if out_path.exists() and not out_path.is_dir():
        raise ValueError(f"Output Folder is not a folder: {out_path}")
    out_path.mkdir(parents=True,exist_ok=True)
    m=builtin_master();rs=rows(bank_path);vs,cnt,total,missing=build(rs,m,auto);o=out_path;stem=bank_path.stem.replace(" ","_")
    a=o/f"{stem}_Voucher_Import.xml";b=o/f"{stem}_Master_Import.xml";c=o/f"{stem}_Voucher_Review.xlsx"
    try:
        voucher_xml(vs,a);master_xml(m,b);review(vs,c)
    except PermissionError as e:
        locked=str(getattr(e,"filename",None) or a)
        raise PermissionError(
            f"Cannot write the output file because Windows denied access.\n\n"
            f"File: {locked}\n\n"
            f"Close that XML/Excel file if it is open in Tally, Excel, Notepad, or another program, "
            f"then click CREATE XML again."
        ) from e
    return len(rs),len(vs),cnt,total,missing,[a,b,c]

class App:
    def __init__(self,r):
        self.r=r;r.title(APP);r.geometry("900x620")
        self.bank=tk.StringVar();self.out=tk.StringVar();self.auto=tk.BooleanVar(value=True)
        ttk.Label(r,text=APP,font=("Segoe UI",20,"bold")).pack(pady=15);f=ttk.Frame(r);f.pack(fill="x",padx=25)
        self.row(f,0,"Bank Statement",self.bank,self.pb);self.row(f,1,"Output Folder",self.out,self.po)
        ttk.Checkbutton(r,text="Auto-create missing ledgers",variable=self.auto).pack(anchor="w",padx=25,pady=8)
        q=ttk.Frame(r);q.pack(pady=5);ttk.Button(q,text="CHECK DATA",command=self.check).pack(side="left",padx=5);ttk.Button(q,text="CREATE XML",command=self.create).pack(side="left",padx=5)
        self.log=tk.Text(r,font=("Consolas",10),state="disabled");self.log.pack(fill="both",expand=True,padx=25,pady=15)
    def row(self,f,i,t,v,cmd):
        ttk.Label(f,text=t,width=18).grid(row=i,column=0,sticky="w",pady=8);ttk.Entry(f,textvariable=v).grid(row=i,column=1,sticky="ew",padx=6);ttk.Button(f,text="Browse",command=cmd).grid(row=i,column=2);f.columnconfigure(1,weight=1)
    def pb(self): 
        p=filedialog.askopenfilename(filetypes=[("Excel","*.xlsx *.xlsm")]);self.bank.set(p) if p else None
    def po(self):
        p=filedialog.askdirectory();self.out.set(p) if p else None
    def write(self,x):
        self.log.config(state="normal");self.log.insert("end",x+"\n");self.log.see("end");self.log.config(state="disabled")
    def check(self):
        try:
            rs=rows(self.bank.get());m=builtin_master();vs,c,t,mi=build(rs,m,False);self.write(f"Transactions: {len(rs)} | Vouchers: {len(vs)}");self.write(f"Journal {c['Journal']} | Payment {c['Payment']} | Contra {c['Contra']} | Receipt {c['Receipt']}");self.write(f"Total: ₹{t:,.2f} | Missing ledgers: {len({k(x[0]) for x in mi})}");messagebox.showinfo("Check",f"{len(rs)} transactions\n{len(vs)} vouchers\nMissing ledgers: {len({k(x[0]) for x in mi})}\nTotal: ₹{t:,.2f}")
        except Exception as e:messagebox.showerror("Error",str(e))
    def create(self):
        try:
            n,v,c,t,mi,fs=generate(self.bank.get(),self.out.get(),self.auto.get());self.write(f"Created {v} vouchers from {n} transactions | ₹{t:,.2f}");[self.write("→ "+str(x)) for x in fs];messagebox.showinfo("Completed",f"Created successfully!\n\nTransactions: {n}\nVouchers: {v}\nTotal: ₹{t:,.2f}")
        except Exception as e:self.write(traceback.format_exc());messagebox.showerror("Error",str(e))
def _fatal_error(exc_type, exc_value, exc_tb):
    try:
        messagebox.showerror(
            APP,
            "The program could not start.\\n\\n" + "".join(traceback.format_exception(exc_type, exc_value, exc_tb))
        )
    except Exception:
        pass

tk.Tk.report_callback_exception=_fatal_error

if __name__=="__main__":
 r=tk.Tk();App(r);r.mainloop()
